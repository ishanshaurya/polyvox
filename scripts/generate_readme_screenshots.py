#!/usr/bin/env python3
"""Generate blurred README screenshots from PolyVox Excel reports.

Usage:
  python scripts/generate_readme_screenshots.py ~/Downloads/CALL_LOG_*.xlsx ...

Or pass a directory:
  python scripts/generate_readme_screenshots.py ~/Downloads
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

import matplotlib.pyplot as plt
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "screenshots"

NAME_HEADERS = re.compile(
    r"agent|name|patient|caller|employee|rep|staff|customer|phone|mobile|email",
    re.I,
)

_REDACT_PATTERNS = [
    (re.compile(r"\+?\d[\d\s\-()]{8,}"), "•••-•••-••••"),
    (re.compile(r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b"), "•••@•••.com"),
    (re.compile(r"Rainbow(?:\s+Children'?s?\s+Hospital)?", re.I), "Acme Health"),
    (re.compile(r"https?://\S+"), "[url]"),
]

FILE_MAP = {
    "CALL_LOG": "call-log.png",
    "AGENT_EVALUATIONS": "agent-evaluations.png",
    "ESCALATIONS": "escalations.png",
    "PATIENT_COMPLAINTS": "patient-complaints.png",
}

MAX_ROWS = 14
MAX_COLS = 10
_agent_counter = 0
_name_counter = 0
_file_counter = 0


def _fake_agent(_: str) -> str:
    global _agent_counter
    _agent_counter += 1
    return f"Agent {chr(64 + ((_agent_counter - 1) % 26) + 1)}"


def _fake_name(_: str) -> str:
    global _name_counter
    _name_counter += 1
    return f"Caller {_name_counter}"


def _fake_file(_: str) -> str:
    global _file_counter
    _file_counter += 1
    buckets = ("b23", "b46", "b7p")
    b = buckets[(_file_counter - 1) % 3]
    return f"call_{_file_counter:03d}_{b}_2026-06-01_a1b2c3d4.mp3"


def _sanitize_text(text: str) -> str:
    for pattern, repl in _REDACT_PATTERNS:
        text = pattern.sub(repl, text)
    if len(text) > 52:
        text = text[:49] + "..."
    return text


def _should_blur_header(header: str) -> str | None:
    h = (header or "").strip().lower()
    if not h:
        return None
    if any(k in h for k in ("file", "stem", "recording", "audio")):
        return "file"
    if "agent" in h:
        return "agent"
    if any(k in h for k in ("patient", "caller", "customer", "name")):
        return "name"
    if any(k in h for k in ("phone", "mobile", "email", "date")):
        return "pii"
    return None


def _cell_value(header: str, value: object, row_idx: int) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if row_idx == 0:
        return text

    blur = _should_blur_header(header)
    if blur == "file":
        return _fake_file(text) if text else ""
    if blur == "agent":
        return _fake_agent(text) if text else ""
    if blur == "name":
        return _fake_name(text) if text else ""
    if blur == "pii":
        if re.match(r"\d{4}-\d{2}-\d{2}", text):
            return "2026-06-01"
        return "•••-•••-••••" if text else ""

    # Filename patterns anywhere (Agent_Name_b23_...)
    if re.search(r"\.mp3$|_b\d+p?_|_[a-f0-9]{8}", text, re.I):
        return _fake_file(text)

    # Heuristic: multi-word capitalized names
    if re.match(r"^[A-Z][a-z]+(\s+[A-Z][a-z]+)+$", text):
        return _fake_agent(text)

    return _sanitize_text(text)


def _sheet_to_table(ws, max_rows: int = MAX_ROWS, max_cols: int = MAX_COLS) -> list[list[str]]:
    rows: list[list[str]] = []
    headers: list[str] = []
    for r_idx, row in enumerate(ws.iter_rows(max_row=max_rows, max_col=max_cols, values_only=True)):
        if r_idx == 0:
            headers = [str(c or "") for c in row]
            rows.append(headers)
            continue
        rows.append([_cell_value(headers[i] if i < len(headers) else "", v, r_idx) for i, v in enumerate(row)])
    return rows or [["(empty)"]]


def _render(table: list[list[str]], title: str, out_path: Path) -> None:
    fig_h = min(9, max(3, 0.38 * len(table) + 1.2))
    fig, ax = plt.subplots(figsize=(15, fig_h))
    ax.axis("off")
    tbl = ax.table(cellText=table, loc="center", cellLoc="left")
    tbl.auto_set_font_size(False)
    tbl.set_fontsize(6.5)
    tbl.scale(1, 1.15)

    # Style header row
    for j in range(len(table[0])):
        cell = tbl[(0, j)]
        cell.set_facecolor("#1e3a5f")
        cell.set_text_props(color="white", fontweight="bold")

    ax.set_title(title, fontsize=12, fontweight="bold", pad=14, color="#1e3a5f")
    fig.patch.set_facecolor("white")
    out_path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(out_path, dpi=160, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    print(f"Wrote {out_path}")


def _detect_type(path: Path) -> str | None:
    upper = path.stem.upper()
    for key in FILE_MAP:
        if key in upper:
            return key
    return None


def process_file(xlsx: Path) -> None:
    global _agent_counter, _name_counter, _file_counter
    kind = _detect_type(xlsx)
    if not kind:
        print(f"Skip unknown file: {xlsx.name}")
        return

    _agent_counter = _name_counter = _file_counter = 0
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb.active
    if ws is None:
        wb.close()
        return

    table = _sheet_to_table(ws)
    title = kind.replace("_", " ").title()
    out_path = OUT / FILE_MAP[kind]
    _render(table, title, out_path)
    wb.close()


def collect_files(args: list[str]) -> list[Path]:
    files: list[Path] = []
    for arg in args:
        p = Path(arg).expanduser()
        if p.is_dir():
            files.extend(sorted(p.glob("*.xlsx")))
        elif p.is_file():
            files.append(p)
        else:
            files.extend(sorted(Path().glob(arg)))
    return files


def main() -> None:
    downloads = Path.home() / "Downloads"
    default_names = [
        "CALL_LOG_20260622_1956.xlsx",
        "AGENT_EVALUATIONS_20260622_1956.xlsx",
        "ESCALATIONS_20260622_1956.xlsx",
        "PATIENT_COMPLAINTS_20260622_1956.xlsx",
    ]

    if len(sys.argv) < 2:
        paths = [downloads / n for n in default_names if (downloads / n).is_file()]
        if not paths:
            print("Usage: python scripts/generate_readme_screenshots.py <file|dir> [...]")
            sys.exit(1)
    else:
        paths = collect_files(sys.argv[1:])

    if not paths:
        print("No xlsx files found.")
        sys.exit(1)

    for p in paths:
        print(f"Processing {p}")
        process_file(p)


if __name__ == "__main__":
    main()

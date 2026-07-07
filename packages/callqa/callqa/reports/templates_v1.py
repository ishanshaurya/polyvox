"""Healthcare cohort — 4 Excel workbook templates."""

from __future__ import annotations

import glob
import json
import os
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

from callqa.config.loader import agent_eval_min_calls
from callqa.state import manifest, metadata
from . import styles as S


def format_qa_note(rec: dict) -> str:
    """Compact QA note for Excel cells — especially ASR-failed calls."""
    note = (rec.get("qa_notes") or "").strip()
    if not note:
        return ""
    if not rec.get("unclear_recording") and rec.get("asr_pass", True):
        return note
    outcome = rec.get("call_outcome") or rec.get("csv_call_outcome") or ""
    disp = rec.get("disposition") or rec.get("csv_disposition") or ""
    disp_short = disp.split("__")[-1].replace("_", " ") if disp else ""
    header = "ASR FAIL — KPIs not scored"
    lines = [header]
    if disp_short and outcome:
        lines.append(f"CSV: {outcome} ({disp_short}) — unverified")
    elif outcome:
        lines.append(f"CSV: {outcome} — unverified")
    # Keep the analyst note to one tight line if it was long prose
    body = note.replace("\n", " ").strip()
    if len(body) > 180:
        body = body[:177] + "..."
    lines.append(body)
    lines.append("Action: re-transcribe or manual review")
    return " | ".join(lines)


def clean_reports(cfg: dict) -> int:
    """
    Delete any existing XLSX files in the reports folder before writing new ones.

    This is critical operational hygiene: reports must represent a single run and
    we never update old sheets in-place.
    """
    out_dir = os.path.join(cfg["output_folder"], "reports")
    if not os.path.isdir(out_dir):
        return 0
    removed = 0
    for path in glob.glob(os.path.join(out_dir, "*.xlsx")):
        try:
            os.remove(path)
            removed += 1
        except OSError:
            pass
    return removed


def _merge_records(cfg: dict, analyses: list[dict]) -> list[dict]:
    """Attach CSV-derived fields from metadata sidecars."""
    out = []
    for rec in analyses:
        stem = Path(rec.get("filename", "")).stem or Path(rec.get("_json_path", "")).stem
        if not stem:
            out.append(rec)
            continue
        meta = metadata.load_sidecar(cfg, stem) or {}
        merged = {**rec}
        merged["call_outcome"] = meta.get("call_outcome") or rec.get("call_outcome", "")
        merged["center"] = meta.get("center", "")
        merged["disposition"] = meta.get("disposition", "")
        merged["asr_pass"] = meta.get("asr_pass", not rec.get("unclear_recording"))
        kpi = rec.get("kpi_scores") or {}
        if any(v not in (None, "") for v in kpi.values()):
            merged["asr_pass"] = True
        merged["_stem"] = stem
        out.append(merged)
    return out


def build_call_log(wb, records, cfg):
    ws = wb.create_sheet("Call Log")
    ws.freeze_panes = "A4"
    kpi_names = [k["name"] for k in cfg["kpis"]]
    headers = (
        ["#", "File", "Agent Name", "Hospital", "Call Date", "Patient Issue",
         "Call Outcome", "Overall Score", "Stars"]
        + kpi_names
        + ["ASR Pass", "Escalation?", "Severity", "QA Notes"]
    )
    col_count = len(headers)
    S.write_title(ws, "CALL CENTER QA — CALL LOG",
                  f"Generated: {datetime.now():%d %b %Y %H:%M}  |  Calls: {len(records)}",
                  col_count)
    S.write_header(ws, headers, 3)

    for i, rec in enumerate(records, 1):
        row = 3 + i
        kpi_scores = rec.get("kpi_scores") or {}
        overall = rec.get("overall_score", "")
        row_data = (
            [i, rec.get("filename", ""), rec.get("agent_name", ""),
             rec.get("hospital", ""), rec.get("call_date", ""),
             rec.get("patient_issue", "") if rec.get("asr_pass") else "",
             rec.get("call_outcome", ""), overall,
             S.stars_str(rec.get("star_rating", 1)) if overall not in (None, "") else ""]
            + [kpi_scores.get(k, "") if rec.get("asr_pass") else "" for k in kpi_names]
            + [
                "Yes" if rec.get("asr_pass") else "No",
                "YES" if rec.get("escalation_required") else "No",
                rec.get("escalation_severity", "") if rec.get("escalation_required") else "",
                format_qa_note(rec),
            ]
        )
        for ci, val in enumerate(row_data, 1):
            cell = ws.cell(row=row, column=ci, value=val)
            if ci == col_count and val:  # QA Notes — wrap long text
                cell.alignment = S.left()
        S.style_row(ws, row, col_count, alternate=(i % 2 == 0))
        if row_data[-1]:
            ws.row_dimensions[row].height = 48

    # Agent Summary sheet
    ws2 = wb.create_sheet("Agent Summary")
    S.write_header(ws2, ["Agent", "Calls", "Avg Score", "Escalations"], 1, bg=S.C_BLUE)
    by_agent = defaultdict(list)
    for rec in records:
        if rec.get("overall_score") not in (None, ""):
            try:
                by_agent[rec.get("agent_name", "")].append(float(rec["overall_score"]))
            except (TypeError, ValueError):
                pass
    for ri, (agent, scores) in enumerate(sorted(by_agent.items()), 2):
        avg = round(sum(scores) / len(scores), 1) if scores else ""
        esc = sum(1 for r in records if r.get("agent_name") == agent and r.get("escalation_required"))
        for ci, val in enumerate([agent, len(scores), avg, esc], 1):
            ws2.cell(row=ri, column=ci, value=val)


def build_escalations(wb, records, cfg):
    ws = wb.create_sheet("Escalations")
    headers = ["#", "File", "Agent", "Date", "Severity", "Reason", "Patient Issue", "QA Notes"]
    escalated = [
        r for r in records
        if r.get("escalation_required") and r.get("asr_pass")
    ]
    S.write_title(ws, "ESCALATION REGISTER", f"Total: {len(escalated)}", len(headers))
    S.write_header(ws, headers, 3, bg=S.C_RED_DEEP)
    for i, rec in enumerate(escalated, 1):
        row = 3 + i
        row_data = [
            i, rec.get("filename"), rec.get("agent_name"), rec.get("call_date"),
            rec.get("escalation_severity"), rec.get("escalation_reason"),
            rec.get("patient_issue"), rec.get("qa_notes"),
        ]
        for ci, val in enumerate(row_data, 1):
            ws.cell(row=row, column=ci, value=val)
        S.style_row(ws, row, len(headers), i % 2 == 0)


def build_complaints(wb, records, cfg):
    ws = wb.create_sheet("Patient Complaints")
    headers = ["#", "Agent", "File", "Date", "Complaint", "Call Outcome", "Score"]
    complaints = []
    for rec in records:
        if not rec.get("asr_pass"):
            continue
        for c in rec.get("patient_complaints") or []:
            if c and str(c).strip():
                complaints.append((rec, str(c).strip()))
    S.write_title(ws, "PATIENT COMPLAINTS", f"Total: {len(complaints)}", len(headers))
    S.write_header(ws, headers, 3, bg=S.C_ORANGE)
    for i, (rec, text) in enumerate(complaints, 1):
        row = 3 + i
        row_data = [
            i, rec.get("agent_name"), rec.get("filename"), rec.get("call_date"),
            text, rec.get("call_outcome"), rec.get("overall_score", ""),
        ]
        for ci, val in enumerate(row_data, 1):
            ws.cell(row=row, column=ci, value=val)
        S.style_row(ws, row, len(headers), i % 2 == 0)


def build_agent_evaluations(wb, records, cfg, min_calls: int = 3):
    ws = wb.create_sheet("Agent Evaluations")
    kpi_names = [k["name"] for k in cfg["kpis"]]
    headers = ["Agent", "Calls Analyzed", "Avg Score", "Stars", "Status"] + [
        f"Avg {k}" for k in kpi_names
    ]
    S.write_title(ws, "AGENT EVALUATIONS", f"Min calls for rollup: {min_calls}", len(headers))
    S.write_header(ws, headers, 3)

    by_agent = defaultdict(list)
    for rec in records:
        if rec.get("asr_pass") and rec.get("overall_score") not in (None, ""):
            by_agent[rec.get("agent_name", "Unknown")].append(rec)

    row = 4
    for agent in sorted(by_agent.keys()):
        calls = by_agent[agent]
        n = len(calls)
        if n < min_calls:
            status = "Under analysis"
            avg = ""
            stars = ""
            kpi_avgs = [""] * len(kpi_names)
        else:
            status = "Complete"
            scores = [float(c["overall_score"]) for c in calls]
            avg = round(sum(scores) / n, 1)
            stars = S.stars_str(round(avg / 20))
            kpi_avgs = []
            for k in kpi_names:
                vals = []
                for c in calls:
                    v = (c.get("kpi_scores") or {}).get(k)
                    if v not in (None, ""):
                        vals.append(float(v))
                kpi_avgs.append(round(sum(vals) / len(vals), 1) if vals else "")

        row_data = [agent, n, avg, stars, status] + kpi_avgs
        for ci, val in enumerate(row_data, 1):
            ws.cell(row=row, column=ci, value=val)
        S.style_row(ws, row, len(headers), row % 2 == 0)
        row += 1


def generate(cfg: dict, batch: int | None = None, *, batch_only: bool = False) -> list[str]:
    analysis_dir = os.path.join(cfg["output_folder"], "analysis")

    records = []
    for path in sorted(glob.glob(os.path.join(analysis_dir, "*.json"))):
        if path.endswith("_error.txt"):
            continue
        with open(path, encoding="utf-8") as f:
            data = json.load(f)
            data["_json_path"] = path
            records.append(data)

    if batch is not None:
        mdata = manifest.load(cfg)
        if batch_only:
            allowed = set(manifest.stems_for_batch(mdata, batch))
        else:
            allowed = set()
            for b in range(1, batch + 1):
                allowed.update(manifest.stems_for_batch(mdata, b))
        records = [
            r for r in records
            if Path(r.get("filename", "")).stem in allowed
            or Path(r.get("_json_path", "")).stem in allowed
        ]

    records = _merge_records(cfg, records)
    if not records:
        print("\n  No analysis results for report.\n")
        return []

    removed = clean_reports(cfg)
    if removed:
        print(f"\n  Cleaned reports folder: removed {removed} old .xlsx file(s)\n")

    ts = datetime.now().strftime("%Y%m%d_%H%M")
    out_dir = os.path.join(cfg["output_folder"], "reports")
    os.makedirs(out_dir, exist_ok=True)

    builders = [
        ("CALL_LOG", build_call_log),
        ("ESCALATIONS", build_escalations),
        ("PATIENT_COMPLAINTS", build_complaints),
        ("AGENT_EVALUATIONS", build_agent_evaluations),
    ]

    min_calls = agent_eval_min_calls(cfg)
    paths = []
    scope = f"batch {batch}" + (" only" if batch_only else " cumulative") if batch else "all"
    print(f"\n  Generating {len(builders)} template reports ({len(records)} calls, {scope})...\n")
    for name, builder in builders:
        wb = Workbook()
        wb.remove(wb.active)
        if name == "AGENT_EVALUATIONS":
            builder(wb, records, cfg, min_calls=min_calls)
        else:
            builder(wb, records, cfg)
        path = os.path.join(out_dir, f"{name}_{ts}.xlsx")
        wb.save(path)
        paths.append(path)
        print(f"  ✓ {name} → {os.path.basename(path)}")

    # Update manifest status
    mdata = manifest.load(cfg)
    for rec in records:
        stem = rec.get("_stem") or Path(rec.get("filename", "")).stem
        if stem and manifest.find_call(mdata, stem):
            manifest.update_call(mdata, stem, status="reported")
    manifest.save(cfg, mdata)

    print(f"\n  Reports saved to: {os.path.abspath(out_dir)}/\n")
    return paths

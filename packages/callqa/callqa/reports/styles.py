"""Shared Excel styling for report builders."""

from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

C_NAVY = "1F3864"
C_BLUE = "2E75B6"
C_RED_DEEP = "C00000"
C_ORANGE = "FF6600"
C_GREEN_DARK = "375623"
C_GREEN_LIGHT = "E2EFDA"
C_RED_BAD = "9C0006"
C_RED_LIGHT = "FFC7CE"
C_WHITE = "FFFFFF"
C_GREY = "F2F2F2"
C_GOLD = "FFD700"


def hdr_font(size=10, color=C_WHITE):
    return Font(name="Arial", bold=True, size=size, color=color)


def cell_font(size=9, bold=False, color="000000"):
    return Font(name="Arial", size=size, bold=bold, color=color)


def fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)


def center():
    return Alignment(horizontal="center", vertical="center", wrap_text=True)


def left():
    return Alignment(horizontal="left", vertical="center", wrap_text=True)


def thin_border():
    s = Side(style="thin", color="CCCCCC")
    return Border(left=s, right=s, top=s, bottom=s)


def stars_str(n):
    n = max(1, min(5, int(n or 1)))
    return "★" * n + "☆" * (5 - n)


def score_color(score):
    if score is None:
        return None
    try:
        s = float(score)
    except (ValueError, TypeError):
        return None
    if s >= 80:
        return (C_GREEN_LIGHT, C_GREEN_DARK)
    if s >= 60:
        return ("FFFF99", "7B6000")
    return (C_RED_LIGHT, C_RED_BAD)


def write_title(ws, title, subtitle, col_count):
    ws.merge_cells(f"A1:{get_column_letter(col_count)}1")
    ws["A1"] = title
    ws["A1"].font = Font(name="Arial", bold=True, size=14, color=C_WHITE)
    ws["A1"].fill = fill(C_NAVY)
    ws["A1"].alignment = center()
    ws.row_dimensions[1].height = 28
    ws.merge_cells(f"A2:{get_column_letter(col_count)}2")
    ws["A2"] = subtitle
    ws["A2"].font = Font(name="Arial", size=9, color="CCCCCC", italic=True)
    ws["A2"].fill = fill(C_NAVY)
    ws["A2"].alignment = center()


def write_header(ws, headers, row, bg=C_NAVY):
    for ci, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=ci, value=h)
        c.font = hdr_font()
        c.fill = fill(bg)
        c.alignment = center()
        c.border = thin_border()
    ws.row_dimensions[row].height = 36


def style_row(ws, row_num, col_count, alternate=False):
    bg = C_GREY if alternate else C_WHITE
    for col in range(1, col_count + 1):
        c = ws.cell(row=row_num, column=col)
        if not c.font or not c.font.bold:
            c.font = cell_font()
        c.fill = fill(bg)
        c.border = thin_border()
        if c.alignment.horizontal not in ("center",):
            c.alignment = left()


def set_widths(ws, widths):
    for ci, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(ci)].width = w
